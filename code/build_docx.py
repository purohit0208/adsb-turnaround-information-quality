#!/usr/bin/env python
"""
Build the submission .docx files from the markdown sources.

Elsevier's reference-checking service could not process the original PDF submission
("Microsoft Word cannot open the file"), so the revision is supplied in editable format.

Produces:
  MANUSCRIPT_REVISED.docx      anonymised, figures appended with captions
  RESPONSE_TO_REVIEWERS.docx
  TITLE_PAGE.docx              author details and declarations
  HIGHLIGHTS.docx              (the original submission uploaded the title page by mistake)
  COVER_LETTER.docx
"""
from __future__ import annotations
import re
from pathlib import Path
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION

ROOT = Path(__file__).resolve().parents[1]
MS = ROOT / "manuscript"
FIGS = ROOT / "figures"
GREY = RGBColor(0x55, 0x55, 0x55)

FIGURE_ORDER = ["fig1_structure", "fig2_buffer", "fig3_information_quality",
                "fig4_decision_timing", "fig5_net_benefit", "fig6_transfer", "fig7_europe"]


def new_doc(base_pt=11):
    d = Document()
    st = d.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(base_pt)
    st.paragraph_format.space_after = Pt(6)
    st.paragraph_format.line_spacing = 1.15
    for s in d.sections:
        s.left_margin = s.right_margin = Inches(1.0)
        s.top_margin = s.bottom_margin = Inches(1.0)
    return d


def add_runs(par, text):
    """Render **bold**, *italic* and `code` inside a paragraph."""
    for tok in re.split(r"(\*\*.+?\*\*|(?<!\*)\*[^*]+?\*(?!\*)|`[^`]+?`)", text):
        if not tok:
            continue
        if tok.startswith("**") and tok.endswith("**"):
            par.add_run(tok[2:-2]).bold = True
        elif tok.startswith("*") and tok.endswith("*"):
            par.add_run(tok[1:-1]).italic = True
        elif tok.startswith("`") and tok.endswith("`"):
            r = par.add_run(tok[1:-1]); r.font.name = "Consolas"; r.font.size = Pt(9)
        else:
            par.add_run(tok)


def add_table(doc, rows):
    hdr = [c.strip() for c in rows[0].strip().strip("|").split("|")]
    body = [[c.strip() for c in r.strip().strip("|").split("|")] for r in rows[2:]]
    t = doc.add_table(rows=1, cols=len(hdr))
    t.style = "Table Grid"
    for i, h in enumerate(hdr):
        cell = t.rows[0].cells[i]
        cell.text = ""
        add_runs(cell.paragraphs[0], h)
        for r in cell.paragraphs[0].runs:
            r.bold = True
            r.font.size = Pt(8.5)
    for row in body:
        cells = t.add_row().cells
        for i, v in enumerate(row[:len(hdr)]):
            cells[i].text = ""
            add_runs(cells[i].paragraphs[0], v)
            for r in cells[i].paragraphs[0].runs:
                r.font.size = Pt(8.5)
    doc.add_paragraph()


def md_to_docx(md: str, doc: Document, skip_headings=()):
    lines = md.split("\n")
    i = 0
    skipping = False
    while i < len(lines):
        ln = lines[i]
        s = ln.strip()
        if s.startswith("#"):
            lvl = len(s) - len(s.lstrip("#"))
            title = s.lstrip("#").strip()
            skipping = any(k.lower() in title.lower() for k in skip_headings)
            if not skipping:
                h = doc.add_heading(level=min(lvl, 4))
                add_runs(h, title)
            i += 1
            continue
        if skipping:
            i += 1
            continue
        if s.startswith("|") and i + 1 < len(lines) and set(lines[i + 1].strip()) <= set("|-: "):
            block = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                block.append(lines[i])
                i += 1
            add_table(doc, block)
            continue
        if s.startswith(">"):
            block = []
            while i < len(lines) and lines[i].strip().startswith(">"):
                block.append(lines[i].strip().lstrip(">").strip())
                i += 1
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.35)
            p.paragraph_format.space_before = Pt(6)
            add_runs(p, " ".join(block))
            for r in p.runs:
                r.italic = True
                r.font.color.rgb = GREY
            continue
        if s in ("---", "***"):
            i += 1
            continue
        if s.startswith(("- ", "* ")):
            p = doc.add_paragraph(style="List Bullet")
            add_runs(p, s[2:])
            i += 1
            continue
        if re.match(r"^\d+\.\s", s):
            p = doc.add_paragraph(style="List Number")
            add_runs(p, re.sub(r"^\d+\.\s", "", s))
            i += 1
            continue
        if not s:
            i += 1
            continue
        buf = []
        while i < len(lines) and lines[i].strip() and not lines[i].strip().startswith(("#", "|", ">", "- ", "* ")) \
                and lines[i].strip() not in ("---", "***"):
            buf.append(lines[i].strip())
            i += 1
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        add_runs(p, " ".join(buf))
    return doc


def build_manuscript():
    md = (MS / "MANUSCRIPT_REVISED.md").read_text(encoding="utf-8")
    caps = {}
    for m in re.finditer(r"\*\*Figure (\d)\.\*\*\s*(.+?)(?=\n\n|\Z)", md, flags=re.S):
        caps[int(m.group(1))] = re.sub(r"\s+", " ", m.group(2)).strip()
    doc = new_doc(11)
    md_to_docx(md, doc, skip_headings=("Figure captions",))
    doc.add_page_break()
    h = doc.add_heading(level=2)
    add_runs(h, "Figures")
    for n, stem in enumerate(FIGURE_ORDER, start=1):
        img = FIGS / f"{stem}.png"
        if not img.exists():
            raise SystemExit(f"missing figure: {img}")
        doc.add_picture(str(img), width=Inches(6.4))
        doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        p = doc.add_paragraph()
        add_runs(p, f"**Figure {n}.** {caps.get(n, '')}")
        for r in p.runs:
            r.font.size = Pt(9)
        doc.add_paragraph()
    out = MS / "MANUSCRIPT_REVISED.docx"
    doc.save(out)
    return out, len(caps)


def build_simple(md_path, out_name, base_pt=11, skip=()):
    doc = new_doc(base_pt)
    md_to_docx(Path(md_path).read_text(encoding="utf-8"), doc, skip_headings=skip)
    out = MS / out_name
    doc.save(out)
    return out


if __name__ == "__main__":
    m, ncap = build_manuscript()
    print(f"wrote {m.name}  ({ncap} figure captions matched)")
    for src, dst in [("RESPONSE_TO_REVIEWERS.md", "RESPONSE_TO_REVIEWERS.docx"),
                     ("TITLE_PAGE.md", "TITLE_PAGE.docx"),
                     ("HIGHLIGHTS.md", "HIGHLIGHTS.docx"),
                     ("COVER_LETTER.md", "COVER_LETTER.docx")]:
        p = MS / src
        if p.exists():
            print("wrote", build_simple(p, dst).name)
